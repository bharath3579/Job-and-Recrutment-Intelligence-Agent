"""Comprehensive test suite for the Job & Recruitment Intelligence Agent."""

import os
import sys
from pathlib import Path
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

# Add project root to sys.path
BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from agent.database import Base, Company, Job, Outreach, Person, RecruitmentSignal
from agent.connectors.career_site_connector import CareerSiteConnector
from agent.connectors.hunter_connector import HunterConnector
from agent.connectors.public_search_connector import PublicSearchConnector
from agent.services.ai_outreach_engine import AIOutreachEngine
from agent.services.company_scout import CompanyScout
from agent.services.deduplicator import Deduplicator
from agent.services.email_ingest import EmailIngest
from agent.services.excel_sync import ExcelSync
from agent.services.full_scanner import FullScanner
from agent.services.global_career_search import GlobalCareerSearch
from agent.services.outreach_generator import OutreachGenerator
from agent.services.query_builder import QueryBuilder
from agent.services.reporter import Reporter
from agent.services.scorer import Scorer
from agent.services.signal_detector import SignalDetector


def get_test_db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = TestingSessionLocal()
    yield db
    db.close()


def test_query_builder():
    queries = QueryBuilder.build_linkedin_signal_queries(max_queries=4)
    assert len(queries) <= 4
    for q in queries:
        assert "query" in q
        assert "site:linkedin.com/posts" in q["query"]

    scout_q = QueryBuilder.build_company_recruiter_query("Cognizant")
    assert "people_search_query" in scout_q
    assert "Cognizant" in scout_q["people_search_query"]


def test_signal_detector():
    # Test Type B: Active Recruitment Signal
    text_active = "We are hiring 10 Data Engineers for Hyderabad. DM your resume directly."
    sig_type, urgency, evidence, method = SignalDetector.classify_text(text_active)
    assert sig_type == SignalDetector.SIGNAL_TYPES["ACTIVE_RECRUITMENT"]
    assert "hiring" in evidence.lower()

    # Test Type C: Recruitment Event (Walk-in)
    text_walkin = "Walk-in interview this Saturday at Cognizant Gachibowli campus for PySpark developers."
    sig_type, urgency, evidence, method = SignalDetector.classify_text(text_walkin)
    assert sig_type == SignalDetector.SIGNAL_TYPES["RECRUITMENT_EVENT"]
    assert urgency == "Urgent"

    # Test Type D: Campus / Off-Campus
    text_campus = "Announcing off-campus drive for 2025 engineering graduates in Azure Data Engineering."
    sig_type, urgency, evidence, method = SignalDetector.classify_text(text_campus)
    assert sig_type == SignalDetector.SIGNAL_TYPES["CAMPUS_OFFCAMPUS"]

    # Test Type E: Referral Opportunity
    text_ref = "Happy to refer candidates for Data Engineer openings. DM me with Job ID."
    sig_type, urgency, evidence, method = SignalDetector.classify_text(text_ref)
    assert sig_type == SignalDetector.SIGNAL_TYPES["REFERRAL_OPPORTUNITY"]


def test_scorer():
    sample_job = {
        "job_title": "Azure Data Engineer (Databricks & PySpark)",
        "skills": "Databricks, PySpark, Azure Data Factory, Delta Lake, SQL, Python",
        "description": "Building metadata-driven ETL pipelines and migrating SQL Server to Delta Lake.",
        "location": "Hyderabad",
        "freshness_tier": "<24h",
        "experience_required": "0-2 years (2025 graduates welcome)",
        "recruitment_signal": "Active Recruitment Signal",
    }
    score, explanation = Scorer.score_job(sample_job)
    assert score >= 80.0
    assert "Databricks" in explanation
    assert "Hyderabad" in explanation


def test_deduplicator():
    job1 = {
        "company": "Tiger Analytics India Pvt Ltd",
        "job_title": "Azure Data Engineer",
        "location": "Hyderabad",
    }
    job2 = {
        "company": "Tiger Analytics",
        "job_title": "Data Engineer - Azure",
        "location": "Hyderabad, Telangana",
    }
    assert Deduplicator.are_jobs_duplicate(job1, job2) is True

    diff_job = {
        "company": "Microsoft",
        "job_title": "Frontend Developer",
        "location": "Bengaluru",
    }
    assert Deduplicator.are_jobs_duplicate(job1, diff_job) is False


def test_outreach_generator():
    conn_req = OutreachGenerator.generate_connection_request(
        recipient_name="Pooja Reddy",
        company="Cognizant",
        job_title="Data Engineer",
        person_type="Recruiter",
    )
    assert len(conn_req) <= 300
    assert "LNMIIT" in conn_req
    assert "Databricks" in conn_req

    ref_req = OutreachGenerator.generate_referral_request(
        recipient_name="Vikram Iyer",
        company="Accenture",
        job_id="JOB-00101",
        job_title="Azure Data Engineer",
    )
    assert "JOB-00101" in ref_req
    assert "LNMIIT" in ref_req
    assert "Bharath Kumar" in ref_req


def test_mode_1_full_scan(test_db):
    res = FullScanner.run_full_scan(test_db)
    assert res["success"] is True
    assert res["qualified_jobs_count"] > 0
    assert res["new_drafts_count"] > 0

    # Verify jobs were inserted
    jobs = test_db.query(Job).all()
    assert len(jobs) > 0

    # Verify outreach drafts are Pending Approval
    drafts = test_db.query(Outreach).filter(Outreach.status == "Pending Approval").all()
    assert len(drafts) > 0


def test_mode_2_company_scout(test_db):
    res = CompanyScout.scout_company(
        test_db,
        company_name="LTIMindtree",
        reason="Campus drive announced / 2025 off-campus inquiry",
    )
    assert res["success"] is True
    assert res["company"] == "LTIMindtree"
    assert res["people_count"] > 0
    assert res["outreach_count"] > 0

    # Verify company exists in DB
    comp = test_db.query(Company).filter(Company.company == "LTIMindtree").first()
    assert comp is not None


def test_email_ingest(test_db):
    # First seed a job
    job = Job(
        job_id="JOB-99999",
        company="Fractal Analytics",
        job_title="Data Engineer",
        location="Remote",
        application_status="Applied",
    )
    test_db.add(job)
    test_db.commit()

    email_text = """
    Dear Candidate,
    We are pleased to invite you for a Technical Interview for the Data Engineer role at Fractal Analytics.
    Please select a time slot for your discussion.
    """
    res = EmailIngest.process_email_text(test_db, email_text, sender="hr@fractal.ai")
    assert res["success"] is True
    assert res["new_status"] == "Interview"
    assert res["company"] == "Fractal Analytics"

    # Verify job status was updated
    updated_job = test_db.query(Job).filter(Job.job_id == "JOB-99999").first()
    assert updated_job.application_status == "Interview"


def test_excel_sync(test_db, tmp_path):
    # Populate some data
    FullScanner.run_full_scan(test_db)
    target_excel = tmp_path / "test_tracker.xlsx"

    excel_path = ExcelSync.sync_database_to_excel(test_db, target_path=target_excel)
    assert os.path.exists(excel_path)

    wb = openpyxl.load_workbook(excel_path)
    sheet_names = wb.sheetnames
    expected_sheets = ["Jobs", "People", "Companies", "Outreach", "Recruitment Signals", "Dashboard"]
    for s in expected_sheets:
        assert s in sheet_names

    # Check headers on Sheet 1 (Jobs)
    ws_jobs = wb["Jobs"]
    assert ws_jobs.cell(row=1, column=1).value == "Job ID"
    assert ws_jobs.cell(row=1, column=2).value == "Company"


def test_reports(test_db):
    FullScanner.run_full_scan(test_db)
    daily = Reporter.generate_daily_report(test_db)
    assert "Daily Job Intelligence" in daily
    assert "High Priority Opportunities" in daily

    weekly = Reporter.generate_weekly_report(test_db)
    assert "Weekly Recruitment & Market Intelligence Report" in weekly


def test_hunter_connector():
    """Test Hunter.io domain search, email finder, and verification."""
    # 1. Domain Search
    domain_res = HunterConnector.domain_search("Databricks")
    assert domain_res["success"] is True
    assert domain_res["domain"] == "databricks.com"
    assert "pattern" in domain_res

    # 2. Email Finder
    finder_res = HunterConnector.find_email("Alex", "Morgan", "Databricks")
    assert finder_res["success"] is True
    assert "@databricks.com" in finder_res["email"]
    assert "score" in finder_res
    assert "verification_status" in finder_res

    # 3. Email Verifier
    verifier_res = HunterConnector.verify_email("alex.morgan@databricks.com")
    assert verifier_res["success"] is True
    assert verifier_res["deliverable"] is True


def test_career_site_connector():
    """Test official career site and ATS integration (Greenhouse/Lever/Direct)."""
    jobs = CareerSiteConnector.search_company_careers("Databricks", role="Data Engineer", country="United States", limit=3)
    assert len(jobs) > 0
    assert "Databricks" in jobs[0]["company"]
    assert "job_url" in jobs[0]
    assert "source" in jobs[0]


def test_public_search_connector():
    """Test public professional search for Recruiters, Managers, and Referrers."""
    people = PublicSearchConnector.discover_people(
        company="Cognizant",
        role="Data Engineer",
        country="India",
        persona_category="All relevant levels",
    )
    assert len(people) >= 3
    types = [p["person_type"] for p in people]
    assert any("Recruiter" in t for t in types)
    assert any("Manager" in t or "Lead" in t for t in types)
    for p in people:
        assert "linkedin_url" in p
        assert "linkedin.com/search" in p["linkedin_url"]


def test_ai_outreach_engine():
    """Test AI-tailored drafts, anti-spam <300 char limits, and timezone recommendations."""
    # Test Job Opportunity Inquiry to Recruiter
    drafts_recruiter = AIOutreachEngine.generate_drafts(
        recipient_name="Pooja Sharma",
        company="Databricks",
        job_title="Data Engineer",
        person_type="Technical Recruiter",
        country="United States",
        objective="Ask about a job opportunity",
        skills="Azure, Databricks, PySpark",
    )
    assert len(drafts_recruiter["linkedin_note"]) <= 300
    assert "Pooja" in drafts_recruiter["linkedin_note"]
    assert "Databricks" in drafts_recruiter["linkedin_note"]
    assert "subject" in drafts_recruiter
    assert "timezone_recommendation" in drafts_recruiter
    assert "EST" in drafts_recruiter["timezone_recommendation"]

    # Test Referral Request (Clean 1st-person, NO copy-paste referral packet)
    drafts_referral = AIOutreachEngine.generate_drafts(
        recipient_name="Karthik R",
        company="Snowflake",
        job_title="Data Platform Engineer",
        person_type="Senior Data Engineer",
        country="India",
        objective="Ask for a referral",
        skills="Databricks, PySpark, Delta Lake",
    )
    assert len(drafts_referral["linkedin_note"]) <= 300
    assert "referral" in drafts_referral["linkedin_note"].lower()
    assert "Bharath Kumar" in drafts_referral["email_body"]
    assert "IST" in drafts_referral["timezone_recommendation"]


def test_global_career_search(test_db):
    """Test End-to-end Global Career Search pipeline."""
    result = GlobalCareerSearch.execute_search(
        db=test_db,
        company="Stripe",
        role="Data Engineer",
        skills="PySpark, Databricks, SQL, Delta Lake",
        country="United States",
        people_to_find="All relevant levels",
        outreach_objective="Ask about a job opportunity",
    )
    assert result["success"] is True
    assert result["company"] == "Stripe"
    assert len(result["jobs_found"]) > 0
    assert len(result["people_found"]) > 0
    assert len(result["drafts_created"]) > 0

    # Verify database state
    comp = test_db.query(Company).filter(Company.company == "Stripe").first()
    assert comp is not None

    people = test_db.query(Person).filter(Person.company == "Stripe").all()
    assert len(people) > 0
    for p in people:
        assert p.email is not None
        assert p.confidence_score > 0
        assert p.country == "United States"

    outreach = test_db.query(Outreach).filter(Outreach.company == "Stripe").all()
    assert len(outreach) > 0
    for o in outreach:
        assert o.status == "Pending Approval"
        assert o.channel in ["LinkedIn Note", "Email"]

