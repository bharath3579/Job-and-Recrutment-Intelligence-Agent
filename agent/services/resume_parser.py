"""ATS Resume Parser for .docx Word documents using python-docx."""

import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx

from agent.config import CANDIDATE_PROFILE, RESUMES_DIR


class ResumeParser:
    """
    Parses ATS-formatted .docx resume documents:
    Extracts contact info, technical skills, projects, and work highlights.
    """

    @classmethod
    def get_available_resumes(cls) -> List[Path]:
        """Find all .docx resume files in the resumes directory."""
        if not RESUMES_DIR.exists():
            RESUMES_DIR.mkdir(parents=True, exist_ok=True)
        return list(RESUMES_DIR.glob("*.docx"))

    @classmethod
    def parse_docx(cls, file_path: Path) -> Dict[str, Any]:
        """
        Parse a .docx resume file and extract structured ATS data.
        """
        if not file_path.exists():
            return {"success": False, "error": f"File {file_path} not found."}

        doc = docx.Document(str(file_path))

        # Extract all text from paragraphs and tables
        full_text_lines = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                full_text_lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and text not in full_text_lines:
                        full_text_lines.append(text)

        full_text = "\n".join(full_text_lines)

        # 1. Contact Information Extraction
        email_match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", full_text)
        phone_match = re.search(r"(?:\+91[\s\-]?)?[6-9]\d{9}", full_text)
        linkedin_match = re.search(r"https?://(?:www\.)?linkedin\.com/in/[\w\-]+", full_text, re.I)
        github_match = re.search(r"https?://(?:www\.)?github\.com/[\w\-]+", full_text, re.I)

        # 2. Extract Skills
        skills_detected = []
        all_candidate_skills = CANDIDATE_PROFILE["primary_skills"] + CANDIDATE_PROFILE["secondary_skills"]
        for s in all_candidate_skills:
            if re.search(rf"\b{re.escape(s)}\b", full_text, re.I):
                skills_detected.append(s)

        # 3. Detect Sections
        sections = cls._extract_sections(full_text_lines)

        return {
            "success": True,
            "filename": file_path.name,
            "filepath": str(file_path),
            "email": email_match.group(0) if email_match else "",
            "phone": phone_match.group(0) if phone_match else "",
            "linkedin": linkedin_match.group(0) if linkedin_match else "",
            "github": github_match.group(0) if github_match else "",
            "skills": list(set(skills_detected)),
            "sections": sections,
            "summary_snippet": full_text[:400],
            "raw_text": full_text,
        }

    @classmethod
    def create_sample_ats_resume(cls, target_path: Optional[Path] = None) -> Path:
        """
        Creates a clean, industry-standard ATS-friendly .docx resume for Bharath Kumar
        tailored to Azure, Databricks & PySpark Data Engineering.
        """
        out_path = target_path or (RESUMES_DIR / "Bharath_Kumar_Azure_Data_Engineer_ATS.docx")
        doc = docx.Document()

        # Title / Header
        title = doc.add_heading("BHARATH KUMAR", level=0)
        title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_paragraph("Azure Data Engineer | Databricks & PySpark Specialist")
        subtitle.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        p_contact = doc.add_paragraph("Hyderabad, India | +91-9876543210 | bharath.kumar@email.com | linkedin.com/in/bharathkumar-de")
        p_contact.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER

        # Education
        doc.add_heading("EDUCATION", level=1)
        doc.add_paragraph(
            "The LNM Institute of Information Technology (LNMIIT), Jaipur\n"
            "Bachelor of Technology in Computer and Communication Engineering (2021 – 2025)\n"
            "Focus: Data Engineering, Distributed Systems, Database Management Systems"
        )

        # Technical Skills
        doc.add_heading("TECHNICAL SKILLS", level=1)
        doc.add_paragraph(
            "• Cloud & Big Data: Microsoft Azure, Databricks, PySpark, Azure Data Factory (ADF), ADLS Gen2, Delta Lake, Unity Catalog\n"
            "• Data Processing & ETL: Metadata-driven Pipelines, Data Migration (SQL Server to Databricks), Data Reconciliation, ETL/ELT\n"
            "• Languages & Databases: Python, SQL, C++, JavaScript, MySQL, MongoDB\n"
            "• Web & Tools: React, Node.js, Git, Azure DevOps, Machine Learning, Data Structures & Algorithms"
        )

        # Projects / Experience
        doc.add_heading("KEY PROJECTS & EXPERIENCE", level=1)
        p1 = doc.add_paragraph()
        p1.add_run("Enterprise SQL Server to Databricks Lakehouse Migration\n").bold = True
        p1.add_run(
            "• Designed and deployed scalable ELT pipelines using Azure Data Factory (ADF) and PySpark to migrate legacy SQL Server tables to Delta Lake on ADLS Gen2.\n"
            "• Implemented Unity Catalog for fine-grained access control and unified governance across lakehouse assets.\n"
            "• Built automated data reconciliation scripts validating row counts and schema integrity with 99.9% accuracy."
        )

        p2 = doc.add_paragraph()
        p2.add_run("Metadata-Driven Dynamic Ingestion Framework\n").bold = True
        p2.add_run(
            "• Developed a metadata-driven ingestion engine in Azure Databricks capable of handling schema evolution and incremental delta loads.\n"
            "• Optimized PySpark shuffle partitions, reducing execution runtime by 35% on multi-node Databricks clusters."
        )

        doc.save(str(out_path))
        return out_path

    @staticmethod
    def _extract_sections(lines: List[str]) -> Dict[str, List[str]]:
        sections = {}
        curr_section = "Header"
        sections[curr_section] = []

        headers = ["education", "skills", "technical skills", "experience", "projects", "certifications", "summary"]
        for line in lines:
            line_clean = line.strip().lower()
            matched = False
            for h in headers:
                if line_clean == h or line_clean.startswith(f"{h}:"):
                    curr_section = h.title()
                    sections[curr_section] = []
                    matched = True
                    break
            if not matched and line.strip():
                sections[curr_section].append(line.strip())

        return sections
